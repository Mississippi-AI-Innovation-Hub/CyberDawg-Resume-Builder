# docx_render.py — DOCX renderer for resume generation.


import os
from core.config import LOGO_PATH
from core.ai import _calc_tenure
from core.layout_measurements import PAGE_WIDTH, PAGE_HEIGHT, SIDEBAR_LEFT, SIDEBAR_RIGHT, SINGLE_COLUMN, pt_to_twips, balance_sidebar_right_columns

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

#  flexible docx renderer

def _drgb(rgb_list):
    """Convert [r,g,b] floats (0-1) to docx RGBColor."""
    return RGBColor(int(rgb_list[0]*255), int(rgb_list[1]*255), int(rgb_list[2]*255))

def _dhex(rgb_list):
    return '%02x%02x%02x' % (int(rgb_list[0]*255), int(rgb_list[1]*255), int(rgb_list[2]*255))

def _d_add_run(para, text, bold=False, italic=False, size=None, color=None, font='Arial'):
    run = para.add_run(text)
    run.font.name = font
    run.bold = bold; run.italic = italic
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    return run

def _d_spacing(para, before=0, after=0):
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)

def _d_keep_next(para):
    pPr = para._element.get_or_add_pPr()
    pPr.append(OxmlElement('w:keepNext'))

def _d_remove_borders(cell):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcBdr = OxmlElement('w:tcBdr')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement(f'w:{side}'); b.set(qn('w:val'), 'none'); tcBdr.append(b)
    tcPr.append(tcBdr)

def _d_set_bg(cell, hex_color):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color); tcPr.append(shd)

def _d_add_rule(para, color_hex):
    pPr = para._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom); pPr.append(pBdr)

def _d_section_header(doc, title, style, primary_hex, accent_hex, font, font_size):
    if style == 'bar_bg':
        tbl = doc.add_table(rows=1, cols=1)
        tbl.style = 'Table Grid'
        cell = tbl.cell(0, 0)
        _d_set_bg(cell, accent_hex); _d_remove_borders(cell)
        p = cell.paragraphs[0]; _d_spacing(p, 2, 2)
        _d_add_run(p, title, bold=True, size=font_size, color=_drgb_from_hex(primary_hex), font=font)
        _d_keep_next(p)
    elif style == 'bold_rule':
        p = doc.add_paragraph(); _d_spacing(p, 14, 2); _d_keep_next(p)
        _d_add_run(p, title, bold=True, size=font_size, color=_drgb_from_hex(primary_hex), font=font)
        _d_add_rule(p, primary_hex)
    elif style in ('spaced_rules', 'spaced_colored'):
        rule_color = accent_hex if style == 'spaced_rules' else primary_hex
        p = doc.add_paragraph(); _d_spacing(p, 12, 4); _d_keep_next(p)
        _d_add_run(p, '  '.join(title), bold=True, size=font_size,
                   color=_drgb_from_hex(primary_hex), font=font)
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        for side in ['top', 'bottom']:
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4' if style == 'spaced_rules' else '6')
            b.set(qn('w:space'), '2'); b.set(qn('w:color'), rule_color)
            pBdr.append(b)
        pPr.append(pBdr)
    elif style == 'bold_italic':
        p = doc.add_paragraph(); _d_spacing(p, 16, 2); _d_keep_next(p)
        _d_add_run(p, title, bold=True, italic=True, size=font_size,
                   color=_drgb_from_hex(primary_hex), font=font)

def _drgb_from_hex(hex_str):
    r = int(hex_str[0:2], 16); g = int(hex_str[2:4], 16); b = int(hex_str[4:6], 16)
    return RGBColor(r, g, b)

def _d_job_line(doc, title, company, date, primary, secondary, font, tenure=None, tenure_color=None, company_color=None):
    # line 1: title left, date right
    p1 = doc.add_paragraph(); _d_spacing(p1, 10, 0); _d_keep_next(p1)
    pPr = p1._element.get_or_add_pPr()
    tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right'); tab.set(qn('w:pos'), '10080')
    tabs.append(tab); pPr.append(tabs)
    _d_add_run(p1, title, bold=True, size=10, color=primary, font=font)
    p1.add_run('\t').font.size = Pt(9)
    _d_add_run(p1, date, italic=True, size=9, color=secondary, font=font)
    # line 2: company (+ tenure right-aligned if provided)
    p2 = doc.add_paragraph(); _d_spacing(p2, 0, 0); _d_keep_next(p2)
    if tenure:
        pPr2 = p2._element.get_or_add_pPr()
        tabs2 = OxmlElement('w:tabs'); tab2 = OxmlElement('w:tab')
        tab2.set(qn('w:val'), 'right'); tab2.set(qn('w:pos'), '10080')
        tabs2.append(tab2); pPr2.append(tabs2)
    _d_add_run(p2, company, italic=True, size=9, color=company_color or secondary, font=font)
    if tenure:
        p2.add_run('\t').font.size = Pt(8.5)
        _d_add_run(p2, tenure, bold=True, size=8.5, color=tenure_color or primary, font=font)

def _d_edu_line(doc, degree, school, date, primary, secondary, font, school_color=None):
    # line 1: degree left, date right
    p1 = doc.add_paragraph(); _d_spacing(p1, 10, 0); _d_keep_next(p1)
    pPr = p1._element.get_or_add_pPr()
    tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right'); tab.set(qn('w:pos'), '10080')
    tabs.append(tab); pPr.append(tabs)
    _d_add_run(p1, degree, bold=True, size=10, color=primary, font=font)
    if date:
        p1.add_run('\t').font.size = Pt(9)
        _d_add_run(p1, date, italic=True, size=9, color=secondary, font=font)
    # line 2: school
    if school:
        p2 = doc.add_paragraph(); _d_spacing(p2, 0, 0); _d_keep_next(p2)
        _d_add_run(p2, school, italic=True, size=9, color=school_color or secondary, font=font)

def _d_bullet(target, text, color, font, marker='\u2022', is_last=False):
    p = target.add_paragraph()
    _d_spacing(p, 0, 2)
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(18))
    _d_add_run(p, f'{marker}\t{text}', size=9, color=color, font=font)
    if not is_last:
        _d_keep_next(p)

def _d_skills_table(doc, items, font, text_color, cell_bg_hex=None):
    while len(items) % 3 != 0: items.append('')
    col_w = Inches(7.0 / 3)
    tbl = doc.add_table(rows=0, cols=3); tbl.style = 'Normal Table'
    for i in range(0, len(items), 3):
        row = tbl.add_row()
        for j in range(3):
            cell = row.cells[j]; _d_remove_borders(cell)
            if cell_bg_hex: _d_set_bg(cell, cell_bg_hex)
            p = cell.paragraphs[0]; _d_spacing(p, 3, 3)
            txt = items[i+j]
            _d_add_run(p, ('\u2022\u00a0\u00a0' + txt) if txt else '', size=9, color=text_color, font=font)

def _d_footer_logo(doc):
    if not os.path.exists(LOGO_PATH): return
    for section in doc.sections:
        footer = section.footer; footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(0)
        fp.add_run().add_picture(LOGO_PATH, width=Pt(100))


# sidebar docx helpers
def _d_set_cell_width(cell, twips):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcW')): tcPr.remove(old)
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(twips)); tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)

def _d_zero_cell_margins(cell):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcMar')): tcPr.remove(old)
    tcMar = OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
        m = OxmlElement(f'w:{side}')
        m.set(qn('w:w'), '0'); m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    tcPr.append(tcMar)

def _d_navy_shape_xml(hex_color, sb_w_pt=170.0, page_h_pt=792.0, shape_id=99):
    SB_EMU = int(sb_w_pt / 72 * 914400)
    PH_EMU = int(page_h_pt / 72 * 914400)
    return f'''<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
      <w:r><w:drawing>
      <wp:anchor distT="0" distB="0" distL="0" distR="0"
                 simplePos="0" relativeHeight="1" behindDoc="1"
                 locked="0" layoutInCell="0" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
        <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
        <wp:extent cx="{SB_EMU}" cy="{PH_EMU}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/>
        <wp:wrapNone/>
        <wp:docPr id="{shape_id}" name="SidebarBg{shape_id}"/>
        <wp:cNvGraphicFramePr/>
        <a:graphic>
          <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
            <wps:wsp>
              <wps:cNvSpPr><a:spLocks noChangeArrowheads="1"/></wps:cNvSpPr>
              <wps:spPr>
                <a:xfrm><a:off x="0" y="0"/><a:ext cx="{SB_EMU}" cy="{PH_EMU}"/></a:xfrm>
                <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                <a:solidFill><a:srgbClr val="{hex_color}"/></a:solidFill>
                <a:ln><a:noFill/></a:ln>
              </wps:spPr>
              <wps:bodyPr/>
            </wps:wsp>
          </a:graphicData>
        </a:graphic>
      </wp:anchor>
      </w:drawing></w:r>
    </w:p>'''

def _d_sb_section_header(cell, title, accent_hex, font, pad=22.0):
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Pt(pad)
    _d_spacing(p, before=12, after=2); _d_keep_next(p)
    _d_add_run(p, '  '.join(title), bold=True, size=8.5,
               color=_drgb_from_hex(accent_hex), font=font)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), accent_hex)
    pBdr.append(bottom); pPr.append(pBdr)

def _d_sb_text(cell, text, font, color_rgb, bold=False, pad=22.0):
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Pt(pad)
    _d_spacing(p, 0, 0)
    _d_add_run(p, text, bold=bold, size=8.5, color=color_rgb, font=font)

def _d_sb_bullet(cell, text, font, color_rgb, pad=32.0):
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Pt(pad)
    _d_spacing(p, 0, 0)
    _d_add_run(p, '\u2014\u00a0' + text, size=8.5, color=color_rgb, font=font)

def _d_sb_main_header(target, title, primary_hex, font):
    blank = target.add_paragraph()
    _d_spacing(blank, 0, 0)
    _d_keep_next(blank)
    blank.add_run('').font.size = Pt(4)
    p = target.add_paragraph()
    _d_spacing(p, 12, 4)
    _d_keep_next(p)
    _d_add_run(p, title, bold=True, size=9.5,
               color=_drgb_from_hex(primary_hex), font=font)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), '6')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), primary_hex)
    pBdr.append(bot)
    pPr.append(pBdr)

def _d_sb_main_job(target, title, company, date, title_rgb, company_rgb, date_rgb, font, main_tw):
    p1 = target.add_paragraph()
    _d_spacing(p1, 8, 0)
    _d_keep_next(p1)
    _d_add_run(p1, title, bold=True, size=10.6, color=title_rgb, font=font)
    p2 = target.add_paragraph()
    _d_spacing(p2, 0, 0)
    _d_keep_next(p2)
    pPr = p2._element.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), str(main_tw - pt_to_twips(SIDEBAR_LEFT['main_right_margin'])))
    tabs.append(tab)
    pPr.append(tabs)
    _d_add_run(p2, company, italic=True, size=9, color=company_rgb, font=font)
    p2.add_run('\t').font.size = Pt(8.5)
    _d_add_run(p2, date, size=8.5, color=date_rgb, font=font)

def _d_sb_main_bullet(target, text, color_rgb, font, is_last=False):
    p = target.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(18))
    _d_spacing(p, 0, 0)
    _d_add_run(p, '\u2022\t' + text, size=9, color=color_rgb, font=font)
    if not is_last:
        _d_keep_next(p)


def render_docx(layout, content, output_path):
    """Flexible DOCX renderer driven by layout + content dicts."""
    L = layout
    primary_hex   = _dhex(L['primary_color'])
    accent_hex    = _dhex(L.get('accent_color', L['primary_color']))
    secondary_hex = _dhex(L.get('secondary_color', [0.33, 0.33, 0.33]))
    text_hex      = _dhex(L.get('text_color', [0.1, 0.1, 0.1]))
    primary_rgb   = _drgb(L['primary_color'])
    secondary_rgb = _drgb(L.get('secondary_color', [0.33, 0.33, 0.33]))
    text_rgb      = _drgb(L.get('text_color', [0.1, 0.1, 0.1]))
    font          = L.get('font', 'Arial')
    header_style  = L.get('header_style', 'centered_name')
    section_style = L.get('section_header_style', 'bold_rule')
    layout_type   = L.get('layout_type', 'single_column')
    bullet_marker = L.get('bullet_marker', '\u2022')
    section_order = L.get('section_order', ['summary','experience','education','certifications','military','skills'])
    sidebar_sections = L.get('sidebar_sections', ['contact','education','certifications','skills'])

    # force all main-column text to be dark and readable on white background.
    def _is_light(c): return sum(c)/3 > 0.45 or max(c) > 0.7
    if _is_light(L.get('text_color', [0.1,0.1,0.1])):
        text_rgb = _drgb([0.1, 0.1, 0.1]); text_hex = _dhex([0.1, 0.1, 0.1])
    if _is_light(L.get('secondary_color', [0.33,0.33,0.33])):
        secondary_rgb = _drgb([0.33, 0.33, 0.33]); secondary_hex = _dhex([0.33, 0.33, 0.33])
    if _is_light(L.get('primary_color', [0.12,0.28,0.53])):
        primary_rgb = _drgb([0.12, 0.28, 0.53]); primary_hex = _dhex([0.12, 0.28, 0.53])

    name = content.get('name', '')
    contact = content.get('contact', '')
    if isinstance(contact, list): contact_list = contact; contact = '  |  '.join(contact)
    else: contact_list = [c.strip() for c in contact.split('|')] if contact else []
    subtitle = content.get('subtitle') or content.get('job_subtitle', '')

    # sidebar left layout

    if layout_type == 'sidebar_left':
        SB_W_PT = SIDEBAR_LEFT['sidebar_width'] + 7  # slight visual offset for DOCX rendering
        SB_PAD = SIDEBAR_LEFT['sidebar_padding']
        MAIN_LEFT = SIDEBAR_LEFT['main_left']
        PAGE_W_PT = PAGE_WIDTH
        PAGE_H_PT = PAGE_HEIGHT
        MAIN_RIGHT_MARGIN = SIDEBAR_LEFT['main_right_margin']
        _is_bw_docx = sum(L['primary_color']) / 3 < 0.15
        sb_text_rgb = _drgb([0.2, 0.2, 0.2]) if _is_bw_docx else _drgb(L.get('sidebar_text_color', [0.82, 0.86, 0.91]))
        sb_school_rgb = _drgb([0.33, 0.33, 0.33]) if _is_bw_docx else _drgb(L.get('sidebar_school_color', [0.66, 0.74, 0.82]))
        accent_rgb = _drgb([0, 0, 0]) if _is_bw_docx else _drgb(L.get('accent_color', L['primary_color']))
        # company name color in main column — must be dark enough to read on white
        acc = L.get('accent_color', L['primary_color'])
        company_rgb = _drgb(acc) if (sum(acc)/3 <= 0.45 and max(acc) <= 0.7) else primary_rgb
        sb_name_rgb = RGBColor(0, 0, 0) if _is_bw_docx else RGBColor(255, 255, 255)
        sb_degree_rgb = RGBColor(0, 0, 0) if _is_bw_docx else RGBColor(255, 255, 255)
        main_rule_hex = _dhex(L.get('main_rule_color', [0.77, 0.82, 0.87]))

        doc = Document()
        for sec in doc.sections:
            sec.left_margin = Pt(0); sec.right_margin = Pt(0)
            sec.top_margin = Pt(0); sec.bottom_margin = Pt(36)
            sec.footer_distance = Pt(18); sec.header_distance = Pt(0)
            sec._sectPr.append(OxmlElement('w:titlePg'))
        doc.styles['Normal'].paragraph_format.space_before = Pt(0)
        doc.styles['Normal'].paragraph_format.space_after = Pt(0)

        # name split
        name_parts = name.split()
        if len(name_parts) >= 2:
            first_name = ' '.join(name_parts[:-1]); last_name = name_parts[-1]
        else:
            first_name = name; last_name = ''

        # 3-column table: sidebar | gap | main
        SB_TW = pt_to_twips(SB_W_PT)
        MAIN_TW = pt_to_twips(PAGE_W_PT - MAIN_LEFT)
        GAP_TW = pt_to_twips(MAIN_LEFT - SB_W_PT)
        tbl = doc.add_table(rows=1, cols=3); tbl.style = 'Normal Table'
        tblPr = tbl._tbl.find(qn('w:tblPr'))
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr'); tbl._tbl.insert(0, tblPr)
        tblW = OxmlElement('w:tblW')
        tblW.set(qn('w:w'), '12240'); tblW.set(qn('w:type'), 'dxa')
        tblPr.append(tblW)
        tblCM = OxmlElement('w:tblCellMar')
        for side in ['top','left','bottom','right']:
            m = OxmlElement(f'w:{side}')
            m.set(qn('w:w'), '0'); m.set(qn('w:type'), 'dxa')
            tblCM.append(m)
        tblPr.append(tblCM)

        sb_cell = tbl.cell(0, 0); gap_cell = tbl.cell(0, 1); main_cell = tbl.cell(0, 2)
        for cell, tw in [(sb_cell, SB_TW), (gap_cell, GAP_TW), (main_cell, MAIN_TW)]:
            _d_set_cell_width(cell, tw); _d_remove_borders(cell); _d_zero_cell_margins(cell)
        # add right padding to main cell so content doesn't touch paper edge
        _mc_tcPr = main_cell._tc.get_or_add_tcPr()
        _mc_mar = OxmlElement('w:tcMar')
        for side in ['top', 'left', 'bottom']:
            m = OxmlElement(f'w:{side}')
            m.set(qn('w:w'), '0'); m.set(qn('w:type'), 'dxa')
            _mc_mar.append(m)
        _mc_right = OxmlElement('w:right')
        _mc_right.set(qn('w:w'), str(pt_to_twips(MAIN_RIGHT_MARGIN))); _mc_right.set(qn('w:type'), 'dxa')
        _mc_mar.append(_mc_right)
        _mc_tcPr.append(_mc_mar)
        if not _is_bw_docx:
            _d_set_bg(sb_cell, primary_hex)

        # sidebar: name
        fn_p = sb_cell.paragraphs[0]; fn_p.paragraph_format.left_indent = Pt(SB_PAD)
        _d_spacing(fn_p, before=48, after=0)
        _d_add_run(fn_p, first_name, bold=True, size=22, color=sb_name_rgb, font=font)
        if last_name:
            ln_p = sb_cell.add_paragraph(); ln_p.paragraph_format.left_indent = Pt(SB_PAD)
            _d_spacing(ln_p, 0, 10)
            _d_add_run(ln_p, last_name, bold=True, size=22, color=accent_rgb, font=font)

        # sidebar: sections
        for sec in sidebar_sections:
            if sec == 'contact' and contact_list:
                _d_sb_section_header(sb_cell, 'CONTACT', accent_hex, font, SB_PAD)
                for line in contact_list:
                    _d_sb_text(sb_cell, line, font, sb_text_rgb, pad=SB_PAD)
            elif sec == 'education' and content.get('education'):
                _d_sb_section_header(sb_cell, 'EDUCATION', accent_hex, font, SB_PAD)
                for e in content['education']:
                    _d_sb_text(sb_cell, e[0], font, sb_degree_rgb, bold=True, pad=SB_PAD)
                    if len(e) > 1 and isinstance(e[1], str):
                        _d_sb_text(sb_cell, e[1], font, sb_school_rgb, pad=SB_PAD)
            elif sec == 'certifications' and content.get('certifications'):
                _d_sb_section_header(sb_cell, 'CERTIFICATIONS', accent_hex, font, SB_PAD)
                for c in content['certifications']:
                    _d_sb_bullet(sb_cell, c, font, sb_text_rgb, pad=SB_PAD+10)
            elif sec == 'skills' and content.get('skills'):
                _d_sb_section_header(sb_cell, 'SKILLS', accent_hex, font, SB_PAD)
                for s in content['skills']:
                    _d_sb_bullet(sb_cell, s, font, sb_text_rgb, pad=SB_PAD+10)

        gap_cell.paragraphs[0].add_run('')

        # main: subtitle
        sub_p = main_cell.paragraphs[0]; _d_spacing(sub_p, before=28, after=6)
        _d_keep_next(sub_p)
        if subtitle:
            _d_add_run(sub_p, subtitle, bold=True, size=14, color=primary_rgb, font=font)

        # main: content sections (everything not in sidebar)
        main_secs = [s for s in section_order if s not in sidebar_sections]
        for sec_name in main_secs:
            if sec_name == 'summary' and content.get('summary'):
                _d_sb_main_header(main_cell, 'PROFESSIONAL SUMMARY', primary_hex, font)
                p = main_cell.add_paragraph(); _d_spacing(p, 4, 0)
                _d_add_run(p, content['summary'], size=9, color=text_rgb, font=font)
            elif sec_name == 'experience' and content.get('jobs'):
                _d_sb_main_header(main_cell, 'PROFESSIONAL EXPERIENCE', primary_hex, font)
                for job in content['jobs']:
                    title, company, date = job[0], job[1], job[2]
                    bullets = job[-1] if isinstance(job[-1], list) else []
                    _d_sb_main_job(main_cell, title, company, date, text_rgb, secondary_rgb, secondary_rgb, font, MAIN_TW)
                    for i, b in enumerate(bullets):
                        _d_sb_main_bullet(main_cell, b, text_rgb, font, is_last=(i==len(bullets)-1))
            elif sec_name == 'military' and content.get('military'):
                _d_sb_main_header(main_cell, 'MILITARY SERVICE', primary_hex, font)
                for m in content['military']:
                    branch, role, date = m[0], m[1], m[2]
                    bullets = m[-1] if isinstance(m[-1], list) else []
                    _d_sb_main_job(main_cell, branch, role, date, text_rgb, secondary_rgb, secondary_rgb, font, MAIN_TW)
                    for i, b in enumerate(bullets):
                        _d_sb_main_bullet(main_cell, b, text_rgb, font, is_last=(i==len(bullets)-1))
            elif sec_name == 'volunteer' and content.get('volunteer'):
                _d_sb_main_header(main_cell, 'VOLUNTEER WORK', primary_hex, font)
                for v in content['volunteer']:
                    title, org, date = v[0], v[1], v[2]
                    bullets = v[-1] if isinstance(v[-1], list) else []
                    _d_sb_main_job(main_cell, title, org, date, text_rgb, secondary_rgb, secondary_rgb, font, MAIN_TW)
                    for i, b in enumerate(bullets):
                        _d_sb_main_bullet(main_cell, b, text_rgb, font, is_last=(i==len(bullets)-1))
            elif sec_name == 'references' and content.get('references'):
                _d_sb_main_header(main_cell, 'REFERENCES', primary_hex, font)
                for r in content['references']:
                    p = main_cell.add_paragraph(); _d_spacing(p, 2, 4)
                    name_r = r[0] if len(r) > 0 else ''
                    _d_add_run(p, name_r, bold=True, size=9, color=text_rgb, font=font)
                    rest = [r[i] for i in range(1, len(r)) if r[i]]
                    if rest:
                        _d_add_run(p, '  |  ' + '  |  '.join(rest), size=9, color=secondary_rgb, font=font)

        # footer + sidebar background
        for section in doc.sections:
            for footer in [section.footer, section.first_page_footer]:
                footer.is_linked_to_previous = False
                for p in list(footer.paragraphs): p._element.getparent().remove(p._element)
                if not _is_bw_docx:
                    footer._element.append(etree.fromstring(
                        _d_navy_shape_xml(primary_hex, SB_W_PT, PAGE_H_PT)))
                if os.path.exists(LOGO_PATH):
                    lp = footer.add_paragraph()
                    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    lp.paragraph_format.space_before = Pt(4)
                    lp.add_run().add_picture(LOGO_PATH, width=Pt(100))

        # no-color sidebar: draw a full-page vertical line via a shape in the footer
        if _is_bw_docx:
            LINE_X_PT = SB_W_PT
            LINE_X_EMU = int(LINE_X_PT / 72 * 914400)
            PH_EMU = int(PAGE_H_PT / 72 * 914400)
            line_xml = f'''<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
              <w:r><w:drawing>
              <wp:anchor distT="0" distB="0" distL="0" distR="0"
                         simplePos="0" relativeHeight="2" behindDoc="1"
                         locked="0" layoutInCell="0" allowOverlap="1">
                <wp:simplePos x="0" y="0"/>
                <wp:positionH relativeFrom="page"><wp:posOffset>{LINE_X_EMU}</wp:posOffset></wp:positionH>
                <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
                <wp:extent cx="0" cy="{PH_EMU}"/>
                <wp:effectExtent l="0" t="0" r="0" b="0"/>
                <wp:wrapNone/>
                <wp:docPr id="98" name="SidebarLine98"/>
                <wp:cNvGraphicFramePr/>
                <a:graphic>
                  <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                    <wps:wsp>
                      <wps:cNvCnPr/>
                      <wps:spPr>
                        <a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="{PH_EMU}"/></a:xfrm>
                        <a:prstGeom prst="line"><a:avLst/></a:prstGeom>
                        <a:ln w="9525"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>
                      </wps:spPr>
                      <wps:bodyPr/>
                    </wps:wsp>
                  </a:graphicData>
                </a:graphic>
              </wp:anchor>
              </w:drawing></w:r>
            </w:p>'''
            for section in doc.sections:
                for footer in [section.footer, section.first_page_footer]:
                    footer._element.append(etree.fromstring(line_xml))

        doc.save(output_path)

        return

    # sidebar right docx layout (template 3 style)
    if layout_type == 'sidebar_right':
        _sr_is_bw_d = sum(L['primary_color']) / 3 < 0.15
        if _sr_is_bw_d:
            sr_dark = _dhex([0.169, 0.169, 0.169])
            sr_med_rgb = _drgb([0.29, 0.29, 0.29])
            sr_grey_rgb = _drgb([0.463, 0.463, 0.463])
            sr_light = _dhex([0.733, 0.733, 0.733])
            sr_near_black = _drgb([0.102, 0.102, 0.102])
            sr_text_rgb = _drgb([0.1, 0.1, 0.1])
            sr_secondary_rgb = _drgb([0.33, 0.33, 0.33])
        else:
            sr_dark = primary_hex
            sr_med_rgb = _drgb(L.get('accent_color', L['primary_color']))
            sr_grey_rgb = secondary_rgb
            sr_light = _dhex([0.733, 0.733, 0.733])
            sr_near_black = _drgb([0.102, 0.102, 0.102])
            sr_text_rgb = text_rgb
            sr_secondary_rgb = secondary_rgb

        LEFT_TW = pt_to_twips(SIDEBAR_RIGHT['left_width'])
        GAP_TW = pt_to_twips(SIDEBAR_RIGHT['right_col_x'] - SIDEBAR_RIGHT['divider_x'])
        RIGHT_TW = pt_to_twips(SIDEBAR_RIGHT['right_width'])

        doc = Document()
        for sec in doc.sections:
            sec.left_margin = Pt(SIDEBAR_RIGHT['left_margin']); sec.right_margin = Pt(SIDEBAR_RIGHT['left_margin'])
            sec.top_margin = Pt(0); sec.bottom_margin = Pt(18)
            sec.header_distance = Pt(0)
            sec.footer_distance = Inches(0.3)
            sec._sectPr.append(OxmlElement('w:titlePg'))
        doc.styles['Normal'].paragraph_format.space_before = Pt(0)
        doc.styles['Normal'].paragraph_format.space_after = Pt(0)

        # Put name, contact, rule, and summary in the first-page header only
        for sec in doc.sections:
            header = sec.first_page_header
            header.is_linked_to_previous = False
            # Name
            name_p = header.paragraphs[0]; _d_spacing(name_p, 28, 4)
            _d_add_run(name_p, name, bold=True, size=36, color=_drgb_from_hex(sr_dark), font=font)
            # Contact
            contact_p = header.add_paragraph(); _d_spacing(contact_p, 0, 0)
            _d_add_run(contact_p, contact, size=9.5, color=sr_grey_rgb, font=font)
            # Thick rule
            rule_p = header.add_paragraph(); _d_spacing(rule_p, 2, 0)
            pPr = rule_p._element.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
            bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '18')
            bot.set(qn('w:space'), '1'); bot.set(qn('w:color'), sr_dark)
            pBdr.append(bot); pPr.append(pBdr)
            # Summary
            if content.get('summary'):
                sp = header.add_paragraph(); _d_spacing(sp, 4, 2); _d_keep_next(sp)
                _d_add_run(sp, '  '.join('PROFESSIONAL SUMMARY'), bold=True, size=9, color=_drgb_from_hex(sr_dark), font=font)
                _d_add_rule(sp, sr_light)
                bp = header.add_paragraph(); _d_spacing(bp, 0, 8)
                _d_add_run(bp, content['summary'], size=9, color=sr_text_rgb, font=font)

        # two-column table: left (experience) | right (edu/skills/certs/military)
        tbl = doc.add_table(rows=1, cols=3); tbl.style = 'Normal Table'
        tblPr = tbl._tbl.find(qn('w:tblPr'))
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr'); tbl._tbl.insert(0, tblPr)
        tblW = OxmlElement('w:tblW')
        tblW.set(qn('w:w'), str(LEFT_TW + GAP_TW + RIGHT_TW)); tblW.set(qn('w:type'), 'dxa')
        tblPr.append(tblW)

        left_cell = tbl.cell(0, 0); gap_cell = tbl.cell(0, 1); right_cell = tbl.cell(0, 2)
        for cell, tw in [(left_cell, LEFT_TW), (gap_cell, GAP_TW), (right_cell, RIGHT_TW)]:
            _d_set_cell_width(cell, tw); _d_remove_borders(cell); _d_zero_cell_margins(cell)

        # add right border to left cell (vertical divider)
        tc = left_cell._tc; tcPr = tc.get_or_add_tcPr()
        tcBdr = OxmlElement('w:tcBdr')
        right_b = OxmlElement('w:right')
        right_b.set(qn('w:val'), 'single'); right_b.set(qn('w:sz'), '2')
        right_b.set(qn('w:space'), '0'); right_b.set(qn('w:color'), sr_light)
        tcBdr.append(right_b); tcPr.append(tcBdr)

        gap_cell.paragraphs[0].add_run('')

        # helper: section header in a cell
        def _sr_cell_header(cell, title):
            p = cell.add_paragraph(); _d_spacing(p, 8, 2); _d_keep_next(p)
            _d_add_run(p, '  '.join(title), bold=True, size=9, color=_drgb_from_hex(sr_dark), font=font)
            _d_add_rule(p, sr_light)

        # left cell: experience, volunteer
        left_sections, right_sections = balance_sidebar_right_columns(content)
        first_left = True

        def _sr_left_entry(title_label, items_key, is_job=True):
            """Render a section in the left column (experience/military/volunteer)."""
            nonlocal first_left
            items = content.get(items_key)
            if not items:
                return
            if first_left:
                p = left_cell.paragraphs[0]; _d_spacing(p, 8, 2); _d_keep_next(p)
                _d_add_run(p, '  '.join(title_label), bold=True, size=9,
                           color=_drgb_from_hex(sr_dark), font=font)
                _d_add_rule(p, sr_light)
                first_left = False
            else:
                _sr_cell_header(left_cell, title_label)
            for item in items:
                title_i, org_i, date_i = item[0], item[1], item[2]
                bullets = item[-1] if isinstance(item[-1], list) else []
                p1 = left_cell.add_paragraph(); _d_spacing(p1, 8, 0); _d_keep_next(p1)
                pPr1 = p1._element.get_or_add_pPr()
                tabs1 = OxmlElement('w:tabs'); tab1 = OxmlElement('w:tab')
                tab1.set(qn('w:val'), 'right'); tab1.set(qn('w:pos'), str(LEFT_TW - 200))
                tabs1.append(tab1); pPr1.append(tabs1)
                _d_add_run(p1, title_i, bold=True, size=11, color=sr_near_black, font=font)
                p1.add_run('\t').font.size = Pt(9)
                _d_add_run(p1, date_i, italic=True, size=9, color=sr_secondary_rgb, font=font)
                p2 = left_cell.add_paragraph(); _d_spacing(p2, 1, 0); _d_keep_next(p2)
                _d_add_run(p2, org_i, italic=True, size=9.5, color=sr_secondary_rgb, font=font)
                for i, b in enumerate(bullets):
                    bp = left_cell.add_paragraph()
                    bp.paragraph_format.left_indent = Pt(18)
                    bp.paragraph_format.first_line_indent = Pt(-18)
                    bp.paragraph_format.tab_stops.add_tab_stop(Pt(18))
                    _d_spacing(bp, 0, 0)
                    _d_add_run(bp, '\u25aa\t' + b, size=9, color=sr_text_rgb, font=font)
                    if i < len(bullets) - 1: _d_keep_next(bp)

        for sec_name in left_sections:
            if sec_name == 'experience':
                _sr_left_entry('PROFESSIONAL EXPERIENCE', 'jobs')
            elif sec_name == 'military':
                _sr_left_entry('MILITARY SERVICE', 'military')
            elif sec_name == 'volunteer':
                _sr_left_entry('VOLUNTEER WORK', 'volunteer')
            elif sec_name == 'references' and content.get('references'):
                if first_left:
                    p = left_cell.paragraphs[0]; _d_spacing(p, 8, 2); _d_keep_next(p)
                    _d_add_run(p, '  '.join('REFERENCES'), bold=True, size=9,
                               color=_drgb_from_hex(sr_dark), font=font)
                    _d_add_rule(p, sr_light)
                    first_left = False
                else:
                    _sr_cell_header(left_cell, 'REFERENCES')
                for r in content['references']:
                    rp = left_cell.add_paragraph(); _d_spacing(rp, 2, 2)
                    _d_add_run(rp, r[0], bold=True, size=8.5, color=sr_near_black, font=font)
                    rest = [r[i] for i in range(1, len(r)) if r[i]]
                    if rest:
                        rp2 = left_cell.add_paragraph(); _d_spacing(rp2, 0, 0)
                        _d_add_run(rp2, '  |  '.join(rest), size=8, color=sr_secondary_rgb, font=font)

        if first_left:
            left_cell.paragraphs[0].add_run('')

        first_right = True
        for sec_name in right_sections:
            if sec_name == 'education' and content.get('education'):
                if first_right:
                    p = right_cell.paragraphs[0]; _d_spacing(p, 8, 2); _d_keep_next(p)
                    _d_add_run(p, '  '.join('EDUCATION'), bold=True, size=9,
                               color=_drgb_from_hex(sr_dark), font=font)
                    _d_add_rule(p, sr_light)
                    first_right = False
                else:
                    _sr_cell_header(right_cell, 'EDUCATION')
                for e in content['education']:
                    dp = right_cell.add_paragraph(); _d_spacing(dp, 4, 0)
                    _d_add_run(dp, e[0], bold=True, size=9.5, color=sr_near_black, font=font)
                    if len(e) > 1 and isinstance(e[1], str):
                        sp = right_cell.add_paragraph(); _d_spacing(sp, 0, 0)
                        _d_add_run(sp, e[1], italic=True, size=9, color=sr_med_rgb, font=font)
                    if len(e) > 2 and isinstance(e[2], str):
                        yp = right_cell.add_paragraph(); _d_spacing(yp, 0, 0)
                        _d_add_run(yp, e[2], size=8.5, color=sr_secondary_rgb, font=font)
            elif sec_name == 'certifications' and content.get('certifications'):
                _sr_cell_header(right_cell, 'CERTIFICATIONS'); first_right = False
                for cert in content['certifications']:
                    cp = right_cell.add_paragraph()
                    cp.paragraph_format.left_indent = Pt(18)
                    cp.paragraph_format.first_line_indent = Pt(-18)
                    cp.paragraph_format.tab_stops.add_tab_stop(Pt(18))
                    _d_spacing(cp, 0, 0)
                    _d_add_run(cp, '\u25aa\t' + cert, size=8.5, color=sr_med_rgb, font=font)
            elif sec_name == 'skills' and content.get('skills'):
                _sr_cell_header(right_cell, 'CORE SKILLS'); first_right = False
                for s in content['skills']:
                    sp = right_cell.add_paragraph()
                    sp.paragraph_format.left_indent = Pt(18)
                    sp.paragraph_format.first_line_indent = Pt(-18)
                    sp.paragraph_format.tab_stops.add_tab_stop(Pt(18))
                    _d_spacing(sp, 0, 0)
                    _d_add_run(sp, '\u25aa\t' + s, size=8.5, color=sr_med_rgb, font=font)
            elif sec_name == 'military' and content.get('military'):
                _sr_cell_header(right_cell, 'MILITARY SERVICE'); first_right = False
                for m in content['military']:
                    mp = right_cell.add_paragraph(); _d_spacing(mp, 4, 0)
                    _d_add_run(mp, m[0], bold=True, size=9.5, color=sr_near_black, font=font)
                    if len(m) > 1:
                        rp = right_cell.add_paragraph(); _d_spacing(rp, 0, 0)
                        _d_add_run(rp, m[1], italic=True, size=9, color=sr_secondary_rgb, font=font)
                    if len(m) > 2:
                        dp = right_cell.add_paragraph(); _d_spacing(dp, 0, 0)
                        _d_add_run(dp, m[2], size=8.5, color=sr_secondary_rgb, font=font)
            elif sec_name == 'references' and content.get('references'):
                _sr_cell_header(right_cell, 'REFERENCES'); first_right = False
                for r in content['references']:
                    rp = right_cell.add_paragraph(); _d_spacing(rp, 2, 2)
                    _d_add_run(rp, r[0], bold=True, size=8.5, color=sr_near_black, font=font)
                    rest = [r[i] for i in range(1, len(r)) if r[i]]
                    if rest:
                        rp2 = right_cell.add_paragraph(); _d_spacing(rp2, 0, 0)
                        _d_add_run(rp2, '  |  '.join(rest), size=8, color=sr_secondary_rgb, font=font)
            elif sec_name == 'board' and content.get('board'):
                _sr_cell_header(right_cell, 'BOARD SERVICE'); first_right = False
                for item in content['board']:
                    bp = right_cell.add_paragraph(); _d_spacing(bp, 0, 0)
                    _d_add_run(bp, item, size=8.5, color=sr_secondary_rgb, font=font)

        if first_right:
            right_cell.paragraphs[0].add_run('')

        _d_footer_logo(doc)
        # Also set first-page footer since titlePg is enabled
        if os.path.exists(LOGO_PATH):
            for sec in doc.sections:
                fp_footer = sec.first_page_footer
                fp_footer.is_linked_to_previous = False
                fp = fp_footer.paragraphs[0]
                fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                fp.paragraph_format.space_before = Pt(0)
                fp.add_run().add_picture(LOGO_PATH, width=Pt(100))
        doc.save(output_path)

        return

    # single column layout (templates 1 and 2 style)
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = Pt(SINGLE_COLUMN['left_margin']); sec.right_margin = Pt(SINGLE_COLUMN['right_margin'])
        sec.top_margin = Pt(0) if header_style == 'banner' else Pt(SINGLE_COLUMN['top_margin'])
        sec.bottom_margin = Pt(SINGLE_COLUMN['bottom_margin'])
    doc.styles['Normal'].paragraph_format.space_before = Pt(0)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)

    # header
    if header_style == 'banner':
        # put banner in Word header so it repeats on every page
        for sec in doc.sections:
            sec.header_distance = Pt(0)
            sec.top_margin = Pt(120)
            header = sec.header
            header.is_linked_to_previous = False
            # remove default empty paragraph
            for p in list(header.paragraphs):
                p._element.getparent().remove(p._element)
            # add banner table in header
            tbl = header.add_table(rows=1, cols=1, width=Inches(8.5))
            tbl.style = 'Normal Table'
            tblPr = tbl._tbl.find(qn('w:tblPr'))
            if tblPr is None:
                tblPr = OxmlElement('w:tblPr'); tbl._tbl.insert(0, tblPr)
            tblW = OxmlElement('w:tblW')
            tblW.set(qn('w:w'), '12240'); tblW.set(qn('w:type'), 'dxa')
            tblPr.append(tblW)
            tblInd = OxmlElement('w:tblInd')
            tblInd.set(qn('w:w'), '-1080'); tblInd.set(qn('w:type'), 'dxa')
            tblPr.append(tblInd)
            tblCM = OxmlElement('w:tblCellMar')
            for side in ['top','left','bottom','right']:
                m = OxmlElement(f'w:{side}')
                m.set(qn('w:w'), '0'); m.set(qn('w:type'), 'dxa')
                tblCM.append(m)
            tblPr.append(tblCM)
            cell = tbl.cell(0, 0)
            _d_set_bg(cell, primary_hex); _d_remove_borders(cell)
            _d_set_cell_width(cell, 12240); _d_zero_cell_margins(cell)
            np = cell.paragraphs[0]; _d_spacing(np, 20, 4)
            np.paragraph_format.left_indent = Pt(54)
            _d_add_run(np, name, bold=True, size=36, color=RGBColor(255,255,255), font=font)
            cp = cell.add_paragraph(); _d_spacing(cp, 0, 16)
            cp.paragraph_format.left_indent = Pt(54)
            accent_rgb_c = _drgb(L.get('accent_color', [0.8, 0.9, 0.9]))
            pipe_rgb = _drgb(L.get('banner_pipe_color', L.get('accent_color', [0.5, 0.72, 0.72])))
            if sum(L['primary_color'])/3 < 0.3:
                accent_rgb_c = RGBColor(255, 255, 255)
                pipe_rgb = RGBColor(200, 200, 200)
            contact_parts = contact.split('|')
            for i, part in enumerate(contact_parts):
                if i > 0:
                    _d_add_run(cp, ' | ', size=9, color=pipe_rgb, font=font)
                _d_add_run(cp, part.strip(), size=9, color=accent_rgb_c, font=font)
    else:
        align = WD_ALIGN_PARAGRAPH.CENTER if header_style == 'centered_name' else WD_ALIGN_PARAGRAPH.LEFT
        from docx.enum.text import WD_LINE_SPACING
        name_p = doc.add_paragraph(); name_p.alignment = align; _d_spacing(name_p, 0, 0)
        name_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        name_p.paragraph_format.line_spacing = Pt(26)
        _d_add_run(name_p, name, bold=True, size=24, color=primary_rgb, font=font)
        rule_p = doc.add_paragraph(); _d_spacing(rule_p, 0, 0)
        rule_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        rule_p.paragraph_format.line_spacing = Pt(4)
        _d_add_rule(rule_p, primary_hex)
        if subtitle:
            sub_p = doc.add_paragraph(); sub_p.alignment = align; _d_spacing(sub_p, 2, 2)
            _d_add_run(sub_p, subtitle, bold=True, size=14, color=primary_rgb, font=font)
        contact_p = doc.add_paragraph(); contact_p.alignment = align; _d_spacing(contact_p, 0, 0)
        _d_add_run(contact_p, contact, size=8.5, color=secondary_rgb, font=font)

    def add_section_header(title):
        _d_section_header(doc, title, section_style, primary_hex, accent_hex, font, 10)

    # sections
    for sec_name in section_order:
        if sec_name == 'summary' and content.get('summary'):
            add_section_header('PROFESSIONAL SUMMARY')
            p = doc.add_paragraph(); _d_spacing(p, 0, 0)
            _d_add_run(p, content['summary'], size=9, color=text_rgb, font=font)

        elif sec_name == 'experience' and content.get('jobs'):
            add_section_header('PROFESSIONAL EXPERIENCE')
            for job in content['jobs']:
                title, company, date = job[0], job[1], job[2]
                bullets = job[-1] if isinstance(job[-1], list) else []
                tenure = _calc_tenure(date) if header_style == 'banner' else None
                _accent_rgb = _drgb(L.get('accent_color', L['primary_color'])) if header_style == 'banner' else None
                _d_job_line(doc, title, company, date, text_rgb, secondary_rgb, font,
                           tenure=tenure, tenure_color=primary_rgb, company_color=_accent_rgb)
                for i, b in enumerate(bullets):
                    _d_bullet(doc, b, text_rgb, font, bullet_marker, is_last=(i==len(bullets)-1))

        elif sec_name == 'education' and content.get('education'):
            add_section_header('EDUCATION')
            for e in content['education']:
                degree = e[0]
                school = e[1] if len(e) > 1 and isinstance(e[1], str) else ''
                date = e[2] if len(e) > 2 and isinstance(e[2], str) else None
                _accent_rgb_edu = _drgb(L.get('accent_color', L['primary_color'])) if header_style == 'banner' else None
                _d_edu_line(doc, degree, school, date, text_rgb, secondary_rgb, font, school_color=_accent_rgb_edu)

        elif sec_name == 'certifications' and content.get('certifications'):
            add_section_header('CERTIFICATIONS')
            certs = content['certifications']
            for i, c in enumerate(certs):
                _d_bullet(doc, c, text_rgb, font, bullet_marker, is_last=(i==len(certs)-1))

        elif sec_name == 'military' and content.get('military'):
            add_section_header('MILITARY SERVICE')
            for m in content['military']:
                branch, role, date = m[0], m[1], m[2]
                bullets = m[-1] if isinstance(m[-1], list) else []
                _accent_rgb_mil = _drgb(L.get('accent_color', L['primary_color'])) if header_style == 'banner' else None
                _mil_tenure = _calc_tenure(date) if header_style == 'banner' else None
                _d_job_line(doc, branch, role, date, text_rgb, secondary_rgb, font,
                           company_color=_accent_rgb_mil, tenure=_mil_tenure, tenure_color=primary_rgb)
                for i, b in enumerate(bullets):
                    _d_bullet(doc, b, text_rgb, font, bullet_marker, is_last=(i==len(bullets)-1))

        elif sec_name == 'skills' and content.get('skills'):
            add_section_header('CORE COMPETENCIES')
            cell_bg = _dhex(L['accent_color']) if L.get('skills_cell_bg') else None
            _d_skills_table(doc, list(content['skills']), font, text_rgb, cell_bg)

        elif sec_name == 'board' and content.get('board'):
            add_section_header('BOARD SERVICE & LEADERSHIP')
            for item in content['board']:
                p = doc.add_paragraph(); _d_spacing(p, 0, 0)
                _d_add_run(p, item, size=9, color=text_rgb, font=font)

        elif sec_name == 'volunteer' and content.get('volunteer'):
            add_section_header('VOLUNTEER WORK')
            for v in content['volunteer']:
                title, org, date = v[0], v[1], v[2]
                bullets = v[-1] if isinstance(v[-1], list) else []
                _accent_rgb_vol = _drgb(L.get('accent_color', L['primary_color'])) if header_style == 'banner' else None
                _d_job_line(doc, title, org, date, text_rgb, secondary_rgb, font, company_color=_accent_rgb_vol)
                for i, b in enumerate(bullets):
                    _d_bullet(doc, b, text_rgb, font, bullet_marker, is_last=(i==len(bullets)-1))

        elif sec_name == 'references' and content.get('references'):
            add_section_header('REFERENCES')
            for r in content['references']:
                p = doc.add_paragraph(); _d_spacing(p, 2, 4)
                name_r = r[0] if len(r) > 0 else ''
                _d_add_run(p, name_r, bold=True, size=9, color=text_rgb, font=font)
                rest = [r[i] for i in range(1, len(r)) if r[i]]
                if rest:
                    _d_add_run(p, '  |  ' + '  |  '.join(rest), size=9, color=secondary_rgb, font=font)

    _d_footer_logo(doc)
    doc.save(output_path)

