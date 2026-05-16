# ai.py — AI functions: resume generation, spelling fix, text extraction, PII masking, sanitization.

import json
import pypdf
import io
import time
import re
import base64

from core.config import (
    s3, bedrock, BUCKET, MODEL_ID, COLOR_PALETTES, LAYOUT_MAP,
    INDUSTRY_INSTRUCTIONS, LEVEL_INSTRUCTIONS, DEFAULT_SECTION_ORDER, SECTION_DATA_KEYS
)
from core.prompts import (
    SPELLING_PROMPT, COLOR_INSTRUCTION_TEMPLATE, BW_INSTRUCTION_TEMPLATE,
    RESUME_GENERATION_PROMPT
)


# shared constants 

MONTHS_MAP = {
    'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'may': 5, 'jun': 6,
    'jul': 7, 'aug': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dec': 12
}

DATE_SEPARATORS = [' - ', ' – ', ' to ', '- ', ' -', '-']

BULLET_CHARS = ('\u2022', '\u2013', '\u25aa', '\u25b8', '\u2014', '-', '*')


# PII masking

_PII_PATTERNS = [
    (r'\b\d{3}-\d{2}-\d{4}\b', '[SSN]'),
    (r'\b(?:dob|date of birth|born)[:\s]+\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b', '[DOB]', re.IGNORECASE),
    (r'\b(?:passport(?:\s*#|\s*no\.?|\s*number)?)[:\s]*[A-Z0-9]{6,9}\b', '[PASSPORT]', re.IGNORECASE),
    (r'\b(?:dl|drivers?\s*license|license)[:\s#]*[A-Z0-9]{5,12}\b', '[DL#]', re.IGNORECASE),
    (r'\b(?:4[0-9]{12}(?:[0-9]{3})?|5[1-5][0-9]{14}|3[47][0-9]{13})\b', '[CARD#]'),
]


def mask_pii(text: str):
    """Detect and replace PII patterns. Returns (cleaned_text, list_of_warnings)."""
    warnings = []
    for entry in _PII_PATTERNS:
        pattern = entry[0]
        placeholder = entry[1]
        flags = entry[2] if len(entry) > 2 else 0

        masked, count = re.subn(pattern, placeholder, text, flags=flags)
        if count:
            warnings.append(f"{placeholder} — {count} instance(s) removed")
            text = masked

    return text, warnings


# extract text from uploaded resume

def extract_text_from_upload(uploaded_file) -> str:
    """Extract text from PDF, DOCX, or image files."""
    name = uploaded_file.name.lower()
    raw = uploaded_file.read()

    if name.endswith(".pdf"):
        return _extract_from_pdf(raw)

    if name.endswith(".docx"):
        return _extract_from_docx(raw)

    if name.endswith((".png", ".jpg", ".jpeg", ".webp")):
        return _extract_from_image(name, raw)

    raise ValueError(f"Unsupported file type: {uploaded_file.name}")


def _extract_from_pdf(raw: bytes) -> str:
    """Try pypdf first, fall back to Textract, then OCR via Bedrock."""
    reader = pypdf.PdfReader(io.BytesIO(raw))
    text = "\n".join(p.extract_text() or "" for p in reader.pages).strip()

    if len(text) >= 50:
        return text

    # PDF has no extractable text — try Textract
    try:
        textract_client = __import__('boto3').client("textract", region_name="us-east-1")
        response = textract_client.detect_document_text(Document={"Bytes": raw})
        lines = [block["Text"] for block in response["Blocks"] if block["BlockType"] == "LINE"]
        return "\n".join(lines)
    except Exception:
        pass

    # Last resort: render pages as images and use Bedrock vision
    import fitz
    pages_text = []
    pdf_doc = fitz.open(stream=raw, filetype="pdf")

    for page in pdf_doc:
        pix = page.get_pixmap(dpi=200)
        img_bytes = pix.tobytes("png")
        b64 = base64.standard_b64encode(img_bytes).decode()

        resp = bedrock.invoke_model(
            modelId=MODEL_ID,
            contentType="application/json",
            accept="application/json",
            body=json.dumps({
                "messages": [{"role": "user", "content": [
                    {"image": {"format": "png", "source": {"bytes": b64}}},
                    {"text": "Extract all text from this resume page exactly as written. Return only the raw text, no commentary."}
                ]}],
                "inferenceConfig": {"maxTokens": 2000, "temperature": 0.0}
            }),
        )
        result = json.loads(resp["body"].read())
        pages_text.append(result["output"]["message"]["content"][0]["text"].strip())

    return "\n\n".join(pages_text)


def _extract_from_docx(raw: bytes) -> str:
    """Extract text from DOCX paragraphs and tables."""
    from docx import Document as _Doc
    doc = _Doc(io.BytesIO(raw))
    texts = []

    for p in doc.paragraphs:
        if p.text.strip():
            texts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        texts.append(p.text)

    return "\n".join(texts)


def _extract_from_image(name: str, raw: bytes) -> str:
    """Extract text from image using Bedrock vision."""
    ext_map = {".png": "png", ".jpg": "jpeg", ".jpeg": "jpeg", ".webp": "webp"}
    ext = next(e for e in ext_map if name.endswith(e))
    b64 = base64.standard_b64encode(raw).decode()

    response = bedrock.invoke_model(
        modelId=MODEL_ID,
        contentType="application/json",
        accept="application/json",
        body=json.dumps({
            "messages": [{"role": "user", "content": [
                {"image": {"format": ext_map[ext], "source": {"bytes": b64}}},
                {"text": "Extract all text from this resume image exactly as written. Return only the raw text, no commentary."}
            ]}],
            "inferenceConfig": {"maxTokens": 2000, "temperature": 0.0}
        }),
    )
    result = json.loads(response["body"].read())
    return result["output"]["message"]["content"][0]["text"].strip()


# load S3 context

def load_s3_context():
    """Load training/style reference PDFs from S3."""
    objects = s3.list_objects_v2(Bucket=BUCKET)
    texts = []

    for o in objects.get("Contents", []):
        if o["Key"].endswith(".pdf"):
            obj = s3.get_object(Bucket=BUCKET, Key=o["Key"])
            reader = pypdf.PdfReader(io.BytesIO(obj["Body"].read()))
            texts.append("\n".join(p.extract_text() or "" for p in reader.pages))

    return "\n\n".join(texts)


# fix spelling

def fix_spelling(candidate_info: str) -> str:
    """Use AI to fix spelling/grammar without changing proper nouns."""
    prompt = SPELLING_PROMPT.format(candidate_info=candidate_info)

    try:
        response = bedrock.invoke_model(
            modelId=MODEL_ID,
            contentType="application/json",
            accept="application/json",
            body=json.dumps({
                "messages": [{"role": "user", "content": [{"text": prompt}]}],
                "inferenceConfig": {"maxTokens": 2000, "temperature": 0.0}
            }),
        )
        result = json.loads(response["body"].read())
        return result["output"]["message"]["content"][0]["text"].strip()
    except Exception:
        return candidate_info


# generate resume

def generate_resume(candidate_info: str, context: str, resume_level: str,
                    color_name: str, target_industry: str = "") -> tuple:
    """Generate a complete resume layout + content via Bedrock."""
    candidate_info = fix_spelling(candidate_info)
    level_instr = LEVEL_INSTRUCTIONS.get(resume_level, "")
    palette = COLOR_PALETTES.get(color_name, COLOR_PALETTES["navy"])
    use_color = color_name != "none"
    layout_override = LAYOUT_MAP.get(resume_level, ('single_column', 'left_name', 'bold_rule'))

    # add industry-specific instructions
    if target_industry:
        if target_industry in INDUSTRY_INSTRUCTIONS:
            level_instr += " " + INDUSTRY_INSTRUCTIONS[target_industry]
        else:
            level_instr += (
                f" {target_industry} industry resume. Emphasize relevant terminology, "
                f"certifications, and quantified achievements specific to this field."
            )

    # build color/layout instruction block
    font = layout_override[3] if len(layout_override) > 3 else 'Arial'
    if use_color:
        color_instruction = COLOR_INSTRUCTION_TEMPLATE.format(
            primary=palette['primary'],
            accent=palette['accent'],
            layout_type=layout_override[0],
            header_style=layout_override[1],
            section_header_style=layout_override[2],
            font=font
        )
    else:
        color_instruction = BW_INSTRUCTION_TEMPLATE.format(
            layout_type=layout_override[0],
            header_style=layout_override[1],
            section_header_style=layout_override[2],
            font=font
        )

    # build final prompt
    prompt = RESUME_GENERATION_PROMPT.format(
        resume_level=resume_level,
        level_instr=level_instr,
        color_instruction=color_instruction,
        context=context[:1000],
        candidate_info=candidate_info
    )

    # call Bedrock
    start = time.perf_counter()
    response = bedrock.invoke_model(
        modelId=MODEL_ID,
        contentType="application/json",
        accept="application/json",
        body=json.dumps({
            "messages": [{"role": "user", "content": [{"text": prompt}]}],
            "inferenceConfig": {"maxTokens": 4000, "temperature": 0.1}
        }),
    )
    latency = time.perf_counter() - start

    # parse response
    result = json.loads(response["body"].read())
    usage = result.get("usage", {})
    raw = result["output"]["message"]["content"][0]["text"].strip()

    # clean markdown fencing if present
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    raw = raw.strip()

    # remove trailing commas (invalid JSON)
    raw = re.sub(r',\s*([\]\}])', r'\1', raw)
    data = json.loads(raw)

    # apply layout overrides
    if layout_override:
        lo = data.get('layout', {})
        lo['layout_type'] = layout_override[0]
        if layout_override[1]:
            lo['header_style'] = layout_override[1]
        if layout_override[2]:
            lo['section_header_style'] = layout_override[2]
        if len(layout_override) > 3 and layout_override[3]:
            lo['font'] = layout_override[3]
        data['layout'] = lo

    return data.get("layout", {}), data.get("content", {}), latency, usage


# date parsing helpers

def _split_date_range(date_str):
    """Split a date range string into (start, end) parts."""
    d = date_str.strip().lower()
    for sep in DATE_SEPARATORS:
        if sep in d:
            parts = d.split(sep, 1)
            return parts[0].strip(), parts[-1].strip()
    return d, d


def _parse_year_month(s):
    """Parse a date string into (year, month) tuple."""
    tokens = s.replace(',', ' ').split()
    year = 0
    month = 1

    for t in tokens:
        if t.isdigit() and len(t) == 4:
            year = int(t)
        else:
            for m, v in MONTHS_MAP.items():
                if t.startswith(m):
                    month = v
                    break

    return year, month


# tenure calculation

def _calc_tenure(date_range):
    """Calculate human-readable tenure from a date range string."""
    if not date_range or not isinstance(date_range, str):
        return None

    start_str, end_str = _split_date_range(date_range)

    # parse end date
    if end_str == 'present':
        from datetime import datetime
        ey, em = datetime.now().year, datetime.now().month
    else:
        ey, em = _parse_year_month(end_str)

    # parse start date
    sy, sm = _parse_year_month(start_str)

    if sy == 0 or ey == 0:
        return None

    total_months = (ey - sy) * 12 + (em - sm)
    if total_months < 1:
        return None

    years = total_months // 12
    months = total_months % 12

    if years > 0:
        return f"{years} yr{'s' if years != 1 else ''}"
    else:
        return f"{months} mo{'s' if months != 1 else ''}"


# date ranking for sorting

def _parse_date_rank(date_str):
    """Return a sortable number from a date string. Higher = more recent."""
    if not date_str or not isinstance(date_str, str):
        return 0

    d = date_str.strip().lower()

    # Use the end date for ranking
    for sep in DATE_SEPARATORS:
        if sep in d:
            d = d.split(sep)[-1].strip()
            break

    if d == 'present':
        return 999999

    year, month = _parse_year_month(d)
    return year * 100 + month


# sanitize content

def _strip_bullet(text):
    """Remove leading bullet characters from a string."""
    if isinstance(text, str):
        t = text.lstrip()
        for ch in BULLET_CHARS:
            if t.startswith(ch):
                t = t[len(ch):].lstrip()
                break
        return t
    return text


def _clean_list_section(content, key):
    """Clean a list section: remove nulls, replace None/null strings with empty."""
    if key not in content or not isinstance(content[key], list):
        return

    content[key] = [x for x in content[key] if x is not None]

    for item in (content[key] or []):
        if isinstance(item, list):
            for idx in range(len(item)):
                if item[idx] is None:
                    item[idx] = ""
                if isinstance(item[idx], str) and item[idx].lower() == 'null':
                    item[idx] = ""
            # strip bullets from the last element if it's a list (bullet points)
            if item and isinstance(item[-1], list):
                item[-1] = [_strip_bullet(b) for b in item[-1] if b is not None]

    if not content[key]:
        content[key] = None


def sanitize(content: dict) -> dict:
    """Clean and normalize AI-generated resume content."""
    # clean simple list sections (skills, certs, board)
    for key in ["skills", "certifications", "board"]:
        if key in content and isinstance(content[key], list):
            content[key] = [_strip_bullet(x) for x in content[key] if x is not None]
            if not content[key]:
                content[key] = None

    # title-case skills
    if content.get('skills') and isinstance(content['skills'], list):
        content['skills'] = [
            s.title() if isinstance(s, str) else s
            for s in content['skills']
        ]

    # clean structured sections
    for key in ["jobs", "education", "military", "volunteer"]:
        _clean_list_section(content, key)

    # clean references separately (same logic but different structure)
    if "references" in content and isinstance(content.get("references"), list):
        content["references"] = [x for x in content["references"] if x is not None]
        for item in (content["references"] or []):
            if isinstance(item, list):
                for idx in range(len(item)):
                    if item[idx] is None:
                        item[idx] = ""
                    if isinstance(item[idx], str) and item[idx].lower() == 'null':
                        item[idx] = ""
        if not content["references"]:
            content["references"] = None

    # sort jobs and education by date (most recent first)
    if content.get("jobs") and isinstance(content["jobs"], list):
        content["jobs"] = sorted(
            content["jobs"],
            key=lambda j: _parse_date_rank(j[2] if len(j) > 2 else ''),
            reverse=True
        )
    if content.get("education") and isinstance(content["education"], list):
        content["education"] = sorted(
            content["education"],
            key=lambda e: _parse_date_rank(e[2] if len(e) > 2 else ''),
            reverse=True
        )

    return content


def sanitize_layout(layout: dict, content: dict) -> dict:
    """Filter section_order to only include sections that have data."""
    layout['section_order'] = [
        s for s in DEFAULT_SECTION_ORDER
        if s in SECTION_DATA_KEYS and content.get(SECTION_DATA_KEYS[s])
    ]
    return layout
